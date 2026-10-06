"""Bilingual self-report worksheet. Parsing cells never certifies documentary evidence."""

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re
from zipfile import ZipFile

from email_validator import validate_email, EmailNotValidError
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .intake_schema import ALL_QUESTIONS, VERSION, questions, VISITOR, STUDENT, WORKER
from .types import Check, Fact, Page
from .evidence import Evidence, validate_value
from .store import digest

START_ROW = 6
FORM_ERROR = "文件格式、版本或案件不匹配，请使用本案信息表 / Invalid workbook, version or case; use this case's worksheet"


def read_rows(path):
    # XLSX is a ZIP container. Bound expansion before handing it to the parser.
    with ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 200 or sum(i.file_size for i in entries) > 30_000_000:
            raise ValueError("Workbook expansion limit")
        if any("externallinks" in i.filename.lower() or "vbaproject" in i.filename.lower() for i in entries):
            raise ValueError("External links and macros are not accepted")
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if set(book.sheetnames) != {"Information", "Documents", "Meta"}:
            raise ValueError("Use the supplied workbook")
        sheet = book["Information"]
        if sheet.max_row > 200 or sheet.max_column > 4:
            raise ValueError("Workbook dimensions exceed template")
        for ws in book:
            if ws.max_row > 200 or ws.max_column > 4:
                raise ValueError("Workbook dimensions exceed template")
            for row in ws:
                if any(cell.data_type == "f" for cell in row):
                    raise ValueError("Formula cells are not accepted; paste values")
        meta = {str(a): str(b or "") for a, b in book["Meta"].iter_rows(max_col=2, values_only=True) if a}
        rows = []
        for index, row in enumerate(sheet.iter_rows(min_row=START_ROW, max_col=4, values_only=True), START_ROW):
            key, _, answer, _ = row
            if key is None and answer is None:
                continue
            if isinstance(answer, datetime):
                answer = answer.date().isoformat()
            rows.append((str(key), "" if answer is None else str(answer).strip(), f"Information!C{index}"))
        return meta, rows
    finally:
        book.close()


def read_form(doc):
    doc.kind, doc.language, doc.content_role = "intake", "en", "evidence"
    try:
        _, rows = read_rows(doc.path)
        doc.pages = [Page(number=1, method="form", text="\n".join(f"{cell} {key}: {value}" for key, value, cell in rows))]
    except Exception as exc:
        doc.problems.append(f"Invalid intake workbook: {type(exc).__name__}")
    return doc


def normalize_answer(q, raw):
    if len(raw) > 700 or not raw or raw.casefold() in {"unknown", "null", "n/a", "不清楚", "不知道", "待定"}:
        raise ValueError("请补充明确答案 / Please give a definite answer")
    if raw.casefold() in {"none", "无"} and q.key not in {"other_names", "travel_history_10y"}:
        raise ValueError("此项需要实际信息 / This field needs actual information")
    if re.search(r"\*{2,}|\b[xX]{3,}\b|\byour .+ here\b", raw, re.I):
        raise ValueError("请替换占位符 / Replace the placeholder")
    if raw.startswith(("=", "+=", "@", "http://", "https://")):
        raise ValueError("请填写答案，不要公式或链接 / Enter an answer, not a formula or link")
    if re.search(r"ignore\s+(?:previous\s+|all\s+)?instructions|mark\s+complete|approve\s+all|忽略.*指令|直接.*(?:通过|完成)|标记.*完成", raw, re.I):
        raise ValueError("请填申请事实，不要处理指令 / Give application facts, not processing instructions")
    if q.kind == "bool":
        values = {"yes": "true", "是": "true", "true": "true", "no": "false", "否": "false", "没有": "false", "false": "false"}
        if raw.casefold() not in values:
            raise ValueError("请填写 yes / no（是/否）")
        return values[raw.casefold()]
    if q.kind == "date":
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", raw):
            raise ValueError("请使用 YYYY-MM-DD / Use YYYY-MM-DD")
        return date.fromisoformat(raw).isoformat()
    if q.kind == "number":
        try:
            number = Decimal(raw.replace(",", ""))
            if not number.is_finite() or number < 0:
                raise ValueError("nonnegative")
            return format(number.normalize(), "f")
        except (InvalidOperation, ValueError):
            raise ValueError("请填非负数字 / Enter a non-negative number") from None
    if q.kind == "email":
        try:
            return validate_email(raw, check_deliverability=False).normalized
        except EmailNotValidError:
            raise ValueError("邮箱格式不正确 / Invalid email format") from None
    if q.kind == "phone" and not re.fullmatch(r"\+[1-9][0-9 ()-]{6,20}", raw):
        raise ValueError("请包含国家区号 / Include country code, e.g. +86")
    choices = {
        "marital": {"single", "married", "civil_partner", "unmarried_partner", "divorced", "widowed"},
        "employment": {"employed", "self_employed", "student", "unemployed", "retired"},
        "funding": {"self", "employer", "other"},
    }
    aliases = {"未婚": "single", "已婚": "married", "离异": "divorced", "丧偶": "widowed", "自费": "self", "本人": "self", "雇主": "employer", "其他": "other", "在职": "employed", "自雇": "self_employed", "学生": "student", "无业": "unemployed", "退休": "retired"}
    if q.kind in choices:
        value = aliases.get(raw, raw.casefold())
        if value not in choices[q.kind]:
            raise ValueError(q.hint)
        return value
    if q.kind == "currency":
        if not re.fullmatch(r"[A-Za-z]{3}", raw):
            raise ValueError("请填三位币种 / Use a three-letter currency")
        return raw.upper()
    if q.key in {"nationality", "residence_country", "application_location"}:
        return validate_value(q.key, raw, raw)
    return raw


def apply_form(case, doc, trace):
    try:
        meta, rows = read_rows(doc.path)
        expected = {q.key for q in questions(case.route)}
        keys = [key for key, _, _ in rows]
        if (not case.route or meta.get("version") != VERSION or meta.get("route") != case.route.value
                or meta.get("case_id") not in {"", case.id} or set(keys) != expected or len(keys) != len(expected)):
            raise ValueError("Case, route, version or field IDs do not match")
    except Exception as exc:
        case.intake_errors["workbook"] = FORM_ERROR
        doc.rejected = True
        trace.setdefault("intake", []).append({"document_id": doc.id, "error": str(exc)})
        return
    case.intake_errors.pop("workbook", None)
    old_forms = {d.id for d in case.documents if d.kind == "intake"}
    for key, raw, cell in rows:
        if not raw:
            continue  # Partial resubmission does not erase previously supplied answers.
        # An explicit correction replaces self-reported form values only. Other
        # sources remain active, so a passport or message contradiction still blocks.
        for fact in case.facts:
            if fact.key == key and fact.source_id in old_forms:
                fact.active = False
        try:
            value = normalize_answer(ALL_QUESTIONS[key], raw)
            item = dict(key=key, value=value, source_id=doc.id, page=1, quote=f"{cell} {key}: {raw}")
            fact = Fact(id=digest(item)[:20], **item)
            old = next((f for f in case.facts if f.id == fact.id), None)
            if old:
                old.active = True
            else:
                case.facts.append(fact)
            case.intake_errors.pop(key, None)
            trace.setdefault("intake", []).append({"document_id": doc.id, "cell": cell, "field": key, "value": value})
        except ValueError as exc:
            case.intake_errors[key] = f"{ALL_QUESTIONS[key].label}: {exc}"
    # Derive age only from an unambiguous DOB; never replace a contradictory stated age.
    e = Evidence(case)
    dob = e.day("date_of_birth")
    for fact in case.facts:
        if fact.key == "age" and fact.source_id in old_forms:
            fact.active = False
    if dob:
        reference = e.day("application_date") or date.today()
        age = reference.year - dob.year - ((reference.month, reference.day) < (dob.month, dob.day))
        source = e.facts("date_of_birth")[0]
        if age >= 0:
            item = dict(key="age", value=str(age), source_id=source.source_id, page=source.page, quote=source.quote)
            fact = Fact(id=digest(item)[:20], **item)
            old = next((f for f in case.facts if f.id == fact.id), None)
            if old:
                old.active = True
            else:
                case.facts.append(fact)
            trace.setdefault("intake_derivations", []).append({"field": "age", "dob": dob.isoformat(), "reference": reference.isoformat(), "value": age})


def information_checks(case):
    e, checks = Evidence(case), []
    for key, message in case.intake_errors.items():
        checks.append(Check(id=f"form:{key}", status="fail", source=VERSION, message=message))
    for q in questions(case.route):
        applicable = True if not q.when else None if e.get(q.when[0]) is None else e.get(q.when[0]) in q.when[1]
        value = e.get(q.key)
        status = "not_applicable" if applicable is False else "unknown" if applicable is None or value is None else "pass"
        if value is not None and applicable is True:
            try:
                normalize_answer(q, value)
            except ValueError:
                status = "fail"
        checks.append(Check(id=f"info:{q.key}", status=status, message=q.label, source=q.source,
                            evidence=[f.id for f in e.facts(q.key)]))
    for key in ("date_of_birth", "home_since"):
        value = e.day(key)
        if value and (value > date.today() or (key == "date_of_birth" and (date.today().year - value.year > 120))):
            checks.append(Check(id=f"form:{key}", status="fail", message=f"{ALL_QUESTIONS[key].label}: 日期不合理 / Check this date", source=VERSION))
    if e.day("course_start") and e.day("course_end") and e.day("course_end") <= e.day("course_start"):
        checks.append(Check(id="form:course_end", status="fail", message="课程结束须晚于开始 / Course end must follow start", source=STUDENT))
    if e.get("offences") == "true":
        checks.append(Check(id="offence_scope", status="unknown", human=True, source=VISITOR, message="违法记录说明需要专业核对 / Offence disclosure needs specialist review"))
    if case.route == "student" and e.get("sponsored_last12m") == "true":
        checks.append(Check(id="sponsor_consent", status="pass" if e.get("sponsor_consent_confirmed", {"sponsorship"}) == "true" else "unknown", source=STUDENT, message="请提供官方资助机构书面同意 / Provide the official sponsor's written consent"))
    if case.route == "skilled_worker" and e.get("worker_atas_required") == "true":
        checks.append(Check(id="worker_atas", status="pass" if e.get("atas_reference", {"atas"}) else "unknown", source=WORKER, message="请提供所需 ATAS / Provide the required ATAS"))
    return checks


def write_form(case, path):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = "Information"
    for row in [
        ["申请信息准备表 / Application preparation worksheet"],
        ["这不是 UKVI 正式申请表，未替您提交 / Not the official UKVI application; nothing is submitted"],
        ["填写 C 列并作为附件回复，可用中英文；在线申请答案须用英文 / Fill column C, reply with attachment; online UKVI answers must be English"],
        ["留空表示稍后补充，回传留空不会删除旧答案；不要填公式 / Blanks keep existing answers; no formulas"],
        ["字段 ID / Field ID", "问题 / Question", "您的回答 / Your answer", "提示 / Help"],
    ]:
        sheet.append(row)
    e = Evidence(case)
    for q in questions(case.route):
        hint = q.hint
        if q.when:
            hint += "；仅适用时填写 / Only if " + q.when[0] + " = " + "/".join(q.when[1])
        sheet.append([q.key, q.label, e.get(q.key) or "", hint])
        row = sheet.max_row
        sheet.cell(row, 3).data_type = "s"  # Never turn a customer's answer into a formula.
        sheet.cell(row, 3).number_format = "@"
        sheet.cell(row, 3).fill = PatternFill("solid", fgColor="EAF4FF")
        sheet.cell(row, 3).font = Font(color="1555A2")
        sheet.row_dimensions[row].height = 58
    for row in sheet:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.column_dimensions["A"].hidden = True
    for col, width in {"B": 48, "C": 42, "D": 78}.items():
        sheet.column_dimensions[col].width = width
    for row in range(1, 5):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        sheet.row_dimensions[row].height = 32
    sheet.freeze_panes = "C6"
    sheet.auto_filter.ref = f"B5:D{sheet.max_row}"
    docs = book.create_sheet("Documents")
    docs.append(["材料清单 / Documents", "适用条件 / Conditions", "依据 / Source"])
    docs.append(["护照个人信息页 / Passport details page", "有效、清晰 / Valid and legible", {"visitor": VISITOR, "student": STUDENT, "skilled_worker": WORKER}.get(case.route, VISITOR)])
    items = {
        "visitor": [("旅行、资金和工作支持材料 / Travel, funding and work evidence", "按个人情况，本版覆盖在职自费 / Circumstance dependent; employed self-funded scope")],
        "student": [("CAS / School CAS details", "学校签发 / Issued by school"), ("资金 / Funds", "按国籍、UKVI要求判断；仍须满足资金条件 / Submission depends on nationality and UKVI request"), ("TB、ATAS、官方资助同意 / TB, ATAS, official sponsor consent", "按居住史、课程和资助判断 / Depends on residence, course and sponsorship")],
        "skilled_worker": [("CoS / Employer CoS details", "雇主签发 / Issued by employer"), ("英语、资金、TB、ATAS / English, funds, TB, ATAS", "适用时；资金可由雇主承担 / Conditional; sponsor maintenance may replace funds evidence"), ("无犯罪证明 / Criminal record certificate", "部分教育、医疗等岗位要求；这些岗位本版未覆盖 / Some occupations require this; outside this MVP scope")],
    }
    for label, condition in items.get(case.route, []):
        docs.append([label, condition, {"visitor": VISITOR, "student": STUDENT, "skilled_worker": WORKER}[case.route]])
    docs.append(["认证译文 / Certified translation", "非英语或威尔士语证明 / Evidence not in English or Welsh", VISITOR])
    docs.append(["自填信息不能代替证明 / Self-report does not replace evidence", "清单会随答案变化 / Checklist changes with answers", VERSION])
    for row in docs:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        docs.row_dimensions[row[0].row].height = 65
    for col in ("A", "B", "C"):
        docs.column_dimensions[col].width = 65
    meta = book.create_sheet("Meta")
    for key, value in {"version": VERSION, "case_id": case.id, "route": str(case.route)}.items():
        meta.append([key, value])
    meta.sheet_state = "hidden"
    book.save(path)
    book.close()
    return str(path)
