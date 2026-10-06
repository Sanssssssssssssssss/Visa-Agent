"""Real API checks that personal notes are not substituted for issuer documents."""
import argparse
import hashlib
import json
from pathlib import Path
import time

from openpyxl import load_workbook
from reportlab.pdfgen import canvas
import pypdfium2 as pdfium

from visa_agent.agent import LiveBudget
from visa_agent.intake import START_ROW
from visa_agent.service import VisaService
from visa_agent.store import write_json
from visa_agent.types import CaseEvent, Status

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "datasets/document-quality"


def prepare():
    """Frozen invented inputs; do not imitate seals, logos or security features."""
    if (DATA / "manifest.json").exists():
        return
    DATA.mkdir(parents=True, exist_ok=True)
    values = {
        "passport": ["passport_name: Lin Chen", "passport_number: N12345067", "passport_expiry: 2036-01-01"],
        "bank": ["bank_holder: Lin Chen", "bank_name: Harbor Bank", "bank_currency: GBP", "bank_minimum: 40000",
                 "bank_start: 2026-08-01", "bank_end: 2026-09-28"],
        "employment": ["employment_name: Lin Chen", "employer: Harbor Design Ltd"],
        "personal-notes": ["These are my own notes for the adviser, typed by me.",
            "My name is Lin Chen. My passport number is N12345067.", "My passport expires on 2036-01-01.",
            "I am not sending a copy of the passport in this file."],
    }
    for name, lines in values.items():
        path = DATA / (name + ".pdf")
        pdf = canvas.Canvas(str(path), invariant=1)
        pdf.setFont("Helvetica", 13)
        pdf.drawString(45, 790, "Personal information" if name != "personal-notes" else "My notes")
        pdf.setFont("Helvetica", 11)
        for index, line in enumerate(lines):
            pdf.drawString(45, 750-index*24, line)
        pdf.save()
    with pdfium.PdfDocument(DATA / "passport.pdf") as pdf:
        page = pdf[0]
        bitmap = page.render(scale=2)
        bitmap.to_pil().save(DATA / "passport-notes.jpg")
        bitmap.close()
        page.close()
    answers = json.loads((ROOT / "datasets/intake/visitor-answers.json").read_text(encoding="utf-8"))
    answers.update(applicant_name="Lin Chen", passport_name="Lin Chen", passport_number="N12345067",
                   employer="Harbor Design Ltd", application_date="2026-10-20")
    book = load_workbook(ROOT / "datasets/intake/visitor-blank-zh-en.xlsx")
    for cells in book["Information"].iter_rows(min_row=START_ROW):
        if cells[0].value in answers:
            cells[2].value = answers[cells[0].value]
    book.save(DATA / "completed-information.xlsx")
    book.close()
    write_json(DATA / "manifest.json", {"purpose":"Adversarial invented notes, NOT genuine documents. Files intentionally omit a test watermark to test classification.",
        "files":{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in DATA.iterdir() if p.suffix in {".pdf", ".jpg", ".xlsx"}}})


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--only", nargs="*")
    args = parser.parse_args()
    prepare()
    # Ordinary labels exercise classification beyond the internal-key format guard.
    natural = DATA / "plain-language.pdf"
    if not natural.exists():
        pdf = canvas.Canvas(str(natural), invariant=1)
        pdf.setFont("Helvetica", 14)
        for index, line in enumerate(["Passport", "Name: Lin Chen", "Passport number: N12345067",
                                      "Nationality: China", "Expiry date: 2036-01-01"]):
            pdf.drawString(45, 790-index*30, line)
        pdf.save()
    if args.output.exists():
        raise ValueError("Use a new output folder")
    manifest = json.loads((DATA / "manifest.json").read_text(encoding="utf-8"))
    for name, sha in manifest["files"].items():
        assert hashlib.sha256((DATA / name).read_bytes()).hexdigest() == sha
    plans = {
        "typed-fields": ["passport.pdf", "bank.pdf", "employment.pdf"],
        "self-written-passport": ["personal-notes.pdf", "bank.pdf", "employment.pdf"],
        "screenshot-of-notes": ["passport-notes.jpg", "bank.pdf", "employment.pdf"],
        "ordinary-labels": ["plain-language.pdf", "bank.pdf", "employment.pdf"],
        "public-bank-sample": [str(ROOT / "external-materials/public-images/01-BOC-deposit-redacted.png")],
    }
    write_json(args.output / "experiment.json", {"plans": plans, "expected":"NOT_COMPLETE; no passport evidence from personal notes",
        "test_mode":False, "extra_input_hashes": {natural.name:hashlib.sha256(natural.read_bytes()).hexdigest()},
        "source_hashes":{str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest()
            for p in (ROOT / "src/visa_agent").rglob("*") if p.suffix in {".py", ".md"}}, "manifest":manifest})
    budget = LiveBudget(args.output / "live-budget.sqlite3")
    reports = []
    for name, files in plans.items():
        if args.only and name not in args.only:
            continue
        service = VisaService(args.output / name, "live", hitl=False, application_forms=True, budget=budget)
        service.create_case(name, test_mode=False)
        started = time.monotonic()
        result = service.handle_event(CaseEvent(case_id=name, event_id="1", text="你好，我去英国旅游，表和材料放附件了。",
            attachments=[str(DATA / "completed-information.xlsx"), *(str(DATA / f) for f in files)]))
        case = service.store.get(name)
        passport = {c.id:c.status for c in case.checks if c.id.startswith("passport")}
        passed = (not result.error and result.status == Status.WAIT_USER and bool(passport)
                  and all(status != "pass" for status in passport.values()))
        traces = service.store.traces(name)
        record = {"id":name, "passed":passed, "state":result.status, "error":result.error, "passport_checks":passport,
            "reply":result.reply, "documents":[{"name":d.name,"kind":d.kind,"role":d.content_role,"methods":[p.method for p in d.pages],"problems":d.problems} for d in case.documents],
            "seconds":round(time.monotonic()-started,3),"usage":{key:sum(t.get('usage',{}).get(key,0) for t in traces) for key in ["requests","input_tokens","output_tokens"]}}
        reports.append(record)
        write_json(args.output / "traces" / (name + ".json"), traces)
        write_json(args.output / "summary.json", {"passed":sum(r['passed'] for r in reports),"cases":len(reports),"http_requests":budget.count(),"results":reports})
        print(json.dumps(record,ensure_ascii=False),flush=True)


if __name__ == "__main__":
    main()
